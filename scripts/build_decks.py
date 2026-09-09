"""Genera las dos presentaciones a partir de las imágenes ya producidas.

    python scripts/build_decks.py

* `reports/presentacion_implementacion.pptx` — capturas reales del dashboard
  (`scripts/capture_dashboard.py`).
* `reports/eda_delitos_lima.pptx` — figuras del cuaderno de EDA
  (`notebooks/eda_delitos_lima.ipynb`).

Las diapositivas llevan **poco texto**: un titular, la imagen, y a lo sumo tres
apuntes cortos. Lo que hay que leer está en el cuaderno y en el README; una
diapositiva que se lee no se escucha.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "reports" / "shots"
FIGS = ROOT / "reports" / "eda"
OUT = ROOT / "reports"

# Misma paleta que el dashboard y que las figuras del cuaderno.
INK = RGBColor(0x21, 0x25, 0x29)
MUTED = RGBColor(0x86, 0x8E, 0x96)
CRIME = RGBColor(0xE8, 0x35, 0x4F)
BG = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xF8, 0xF9, 0xFA)

W, H = Inches(13.333), Inches(7.5)
FONT = "Segoe UI"


def _txt(slide, text, x, y, w, h, *, size=18, bold=False, color=INK,
         align=PP_ALIGN.LEFT, space=1.0):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = space
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT
    return box


def _bullets(slide, items, x, y, w, *, size=13):
    box = slide.shapes.add_textbox(x, y, w, Inches(2.4))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (head, rest) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.28
        p.space_after = Pt(7)
        a = p.add_run()
        a.text = head
        a.font.size = Pt(size)
        a.font.bold = True
        a.font.color.rgb = INK
        a.font.name = FONT
        if rest:
            b = p.add_run()
            b.text = "  " + rest
            b.font.size = Pt(size)
            b.font.color.rgb = MUTED
            b.font.name = FONT
    return box


def _bg(slide, color=BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _rule(slide, x, y, w, color=CRIME, h=Emu(28575)):
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def _fit(path: Path, max_w: Emu, max_h: Emu) -> tuple[Emu, Emu]:
    """Tamaño que cabe en la caja conservando la proporción de la imagen."""
    from PIL import Image
    try:
        with Image.open(path) as im:
            iw, ih = im.size
    except Exception:
        return max_w, max_h
    scale = min(max_w / iw, max_h / ih)
    return Emu(int(iw * scale)), Emu(int(ih * scale))


def title_slide(prs, title, subtitle, meta):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s)
    _rule(s, Inches(1.0), Inches(2.5), Inches(1.1))
    _txt(s, title, Inches(1.0), Inches(2.75), Inches(11), Inches(1.4),
         size=40, bold=True, space=1.05)
    _txt(s, subtitle, Inches(1.0), Inches(4.15), Inches(11), Inches(0.9),
         size=17, color=MUTED, space=1.25)
    _txt(s, meta, Inches(1.0), Inches(6.4), Inches(11), Inches(0.5),
         size=11, color=MUTED)
    return s


def section_slide(prs, kicker, title, note=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, SOFT)
    _txt(s, kicker.upper(), Inches(1.0), Inches(2.7), Inches(11), Inches(0.4),
         size=12, bold=True, color=CRIME)
    _txt(s, title, Inches(1.0), Inches(3.15), Inches(11), Inches(1.2),
         size=32, bold=True)
    if note:
        _txt(s, note, Inches(1.0), Inches(4.5), Inches(10), Inches(1.2),
             size=14, color=MUTED, space=1.3)
    return s


def figure_slide(prs, title, image: Path, bullets=None, caption=""):
    """Titular arriba, imagen a la izquierda, apuntes a la derecha.

    Si no hay apuntes la imagen ocupa el ancho completo: una captura grande y
    legible vale más que una pequeña con hueco al lado.
    """
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s)
    _txt(s, title, Inches(0.62), Inches(0.42), Inches(12.1), Inches(0.7),
         size=23, bold=True)
    _rule(s, Inches(0.62), Inches(1.06), Inches(0.62))

    # Se acepta la ruta con o sin extensión: quien escribe el guion de la
    # presentación piensa en nombres de figura, no en ficheros.
    if not image.suffix:
        image = image.with_suffix(".png")
    if not image.exists():
        _txt(s, f"(falta {image.name})", Inches(0.62), Inches(3),
             Inches(6), Inches(0.5), size=13, color=MUTED)
        return s

    if bullets:
        box_w, box_x = Inches(8.5), Inches(0.62)
        w, h = _fit(image, box_w, Inches(5.5))
        s.shapes.add_picture(str(image), box_x, Inches(1.45), width=w, height=h)
        _bullets(s, bullets, Inches(9.4), Inches(1.55), Inches(3.4))
    else:
        w, h = _fit(image, Inches(12.1), Inches(5.5))
        x = Emu(int((W - w) / 2))
        s.shapes.add_picture(str(image), x, Inches(1.45), width=w, height=h)

    if caption:
        _txt(s, caption, Inches(0.62), Inches(7.0), Inches(12.1), Inches(0.4),
             size=10.5, color=MUTED)
    return s


def stat_slide(prs, title, stats, note=""):
    """Cuatro cifras grandes. Para lo que no necesita gráfico."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s)
    _txt(s, title, Inches(0.62), Inches(0.42), Inches(12.1), Inches(0.7),
         size=23, bold=True)
    _rule(s, Inches(0.62), Inches(1.06), Inches(0.62))

    n = len(stats)
    gap, x0, top = Inches(0.35), Inches(0.62), Inches(2.2)
    cell = Emu(int((Inches(12.1) - gap * (n - 1)) / n))
    for i, (num, lab, sub) in enumerate(stats):
        x = Emu(int(x0 + i * (cell + gap)))
        _txt(s, num, x, top, cell, Inches(1.1), size=40, bold=True, color=CRIME)
        _txt(s, lab, x, Inches(3.35), cell, Inches(0.5), size=14, bold=True)
        if sub:
            _txt(s, sub, x, Inches(3.85), cell, Inches(1.2), size=11.5,
                 color=MUTED, space=1.3)
    if note:
        _txt(s, note, Inches(0.62), Inches(5.9), Inches(12.1), Inches(1.0),
             size=13, color=MUTED, space=1.3)
    return s


def new_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    return prs
